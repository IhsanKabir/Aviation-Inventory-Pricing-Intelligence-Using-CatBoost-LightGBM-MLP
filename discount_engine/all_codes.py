"""Every discount code an OTA offered, not just the one the grid shows.

The summary and per-route grids keep ONE common rate and ONE special coupon per
cell (the best). This list keeps all of them: every promo code each B2C channel
served, with its real code name, who can use it, the published rate, what it was
actually worth at the fares seen (caps applied), the cap, the payment fee, whether
it stacks on the automatic discount, the travel dates and the % basis.

Tier:
  * Common       - anyone gets it (automatic discount, wallet or any-payment promo).
  * Special      - needs a specific card, bank or membership (Stellar, AMEX, GPStar,
                   a GoZayaan campaign limited to named banks...).
  * Unclear      - a FirstTrip code that is neither a known card nor a known wallet.
  * No discount  - a 0% code (EMI, internet banking): a payment option, not a rate.
  * Not captured - ShareTrip coupons differ by airline and this airline's booking page
                   wasn't captured, so its coupons are unknown (never borrowed).

Only the coupon channels are listed (FirstTrip B2C, ShareTrip, GoZayaan). The B2B
channels and Amy pay a commission with no promo code; their rates are in the grids.

Rows come from the parses build_report already did (memoized), so nothing is read
or fetched twice. One entry per (route, OTA, airline, tier, code); a code seen on
several fares shows its effective rate as the lowest-highest range, and "Fares seen"
counts DISTINCT fares (date + price), so one fare met twice counts once.
"""
from __future__ import annotations

import math
from typing import Any, Optional

from . import grid as g

COMMON, SPECIAL, UNCLEAR = "Common", "Special", "Unclear"
NO_DISCOUNT, NOT_CAPTURED = "No discount", "Not captured"
# a coupon rate from FirstTrip's coupon table for a route this run did not search: the
# route's automatic discount and its other offers are unknown until it is searched
NOT_SEARCHED = "Not searched"
TIER_ORDER = {COMMON: 0, SPECIAL: 1, UNCLEAR: 2, NOT_CAPTURED: 3, NOT_SEARCHED: 4, NO_DISCOUNT: 5}
BASE, TOTAL = "Base fare", "Booking total"
NO_CODE = "(automatic, no code)"
ALL_ROUTES = "All its routes"     # route of a per-airline row (0% codes)

# Column order for the sheet and the app's table: (key, heading).
COLUMNS = [
    ("market", "Market"), ("route", "Route"), ("travel_dates", "Travel date(s)"),
    ("ota", "OTA"), ("airline", "Airline"), ("tier", "Tier"), ("code", "Promo code"),
    ("who", "Who can use it"), ("published_pct", "Published %"), ("effective", "Effective %"),
    ("basis", "% of"), ("cap_bdt", "Cap (BDT)"), ("stacks", "On top of automatic"),
    ("fee_pct", "Payment fee %"), ("seen", "Fares seen"), ("title", "Offer text"),
]


def _market(flight_type: Any, origin: str = "", destination: str = "") -> str:
    if origin and destination:
        return g._route_type(origin, destination)
    return "DOM" if str(flight_type).upper() == "DOM" else "INTL"


def _route(r: dict[str, Any]) -> str:
    o, d = str(r.get("origin") or "").upper(), str(r.get("destination") or "").upper()
    return f"{o}-{d}" if o and d else ""


def _obs(*, market: str, route: str, ota: str, airline: str, tier: str, code: str,
         who: str, published: Optional[float], effective: Optional[float], basis: str,
         fare: Any = None, date: str = "", cap: Any = None, stacks: Optional[bool] = None,
         fee: Any = None, title: str = "") -> dict[str, Any]:
    """One sighting of one code on one fare (`fare` identifies the fare)."""
    return {"market": market, "route": route, "ota": ota, "airline": airline,
            "tier": tier, "code": code, "who": who,
            "published_pct": published, "effective_pct": effective, "basis": basis,
            "fare": fare, "date": date, "cap_bdt": round(float(cap)) if cap else None,
            "stacks": stacks, "fee_pct": fee, "title": title}


# --- ShareTrip ----------------------------------------------------------------------

def _sharetrip_obs(where: dict[str, Any], auto_pct: float, auto_code: Optional[str],
                   terms: Optional[list[dict[str, Any]]], cell: Optional[dict[str, Any]],
                   titles: dict[str, str]) -> list[dict[str, Any]]:
    """Sightings for one ShareTrip fare: the automatic discount; then, when this airline's
    coupon terms are known, the wallet stack (Common), every other coupon (Special) and
    the 0% codes (No discount); when they aren't, one 'Not captured' marker."""
    # The default coupon (FLIGHTINT / FLYINSIDE) is the one the automatic rate rides on.
    default = next((str(c.get("couponCode")) for c in terms or []
                    if str(c.get("isDefault", "")).lower() in ("1", "true", "yes")), None)
    auto_code = auto_code or default
    # 0% codes are a per-airline payment option, not a per-route rate: one row per
    # airline across all its routes (else EMI/net-banking repeat on every route).
    anywhere = {**where, "route": ALL_ROUTES}
    if auto_pct > 0:
        out = [_obs(**where, tier=COMMON, code=auto_code or NO_CODE, who="Anyone (automatic)",
                    published=auto_pct, effective=auto_pct, title=titles.get(auto_code or "", ""))]
    else:      # e.g. the low-cost carriers' BUDGETFLY: no automatic discount at all; kept on
        # its route so the airline still gets a row in the by-type view
        out = [_obs(**where, tier=NO_DISCOUNT, code=auto_code or NO_CODE,
                    who="Anyone (no automatic discount)", published=0.0, effective=None,
                    title=titles.get(auto_code or "", ""))]
    if terms is None:
        out.append(_obs(**where, tier=NOT_CAPTURED, code="(booking page not captured)",
                        who="Unknown", published=None, effective=None,
                        title="ShareTrip's coupons differ by airline. Open this airline's "
                              "booking page on ShareTrip and save the HAR to see them."))
        return out
    # open to anyone: the wallet coupon and the telco codes (ShareTrip doesn't verify them)
    open_codes = set((cell or {}).get("open_codes") or [])
    telco = {c.get("couponCode") for c in g.sharetrip_har._telco_coupons(terms)}
    for j in (cell or {}).get("judged") or []:
        is_open = j["code"] in open_codes
        who = (f"Anyone entering the {j['label']} code (not verified)" if j["code"] in telco
               else f"{j['label']} payment" if is_open else j["label"])
        out.append(_obs(**where, tier=COMMON if is_open else SPECIAL, code=j["code"], who=who,
                        published=j["nominal_pct"], effective=j["effective_pct"],
                        cap=j.get("cap_bdt"), stacks=j["stacks_with_auto"],
                        fee=j.get("fee_pct"), title=titles.get(j["code"], "")))
    for c in terms:
        code = str(c.get("couponCode") or "")
        if code and code not in (auto_code, default) and float(c.get("discount") or 0) <= 0:
            out.append(_obs(**anywhere, tier=NO_DISCOUNT, code=code,
                            who=g.sharetrip_har._card_label(c), published=0.0, effective=None,
                            stacks=str(c.get("withDiscount", "")).lower() == "yes",
                            title=str(c.get("title") or "")))
    return out


def sharetrip(hars: Optional[list[str]]) -> list[dict[str, Any]]:
    st = g.sharetrip_har
    details: list[dict[str, Any]] = []
    search: list[dict[str, Any]] = []
    for h in hars or []:
        d, s = g._recall("sharetrip_routed", h, ([], []))
        details += d
        search += s
    if not details and not search:
        return []
    gateways = g._recall("sharetrip_gateways", "all", {}) or None
    titles = {str(c.get("couponCode")): str(c.get("title") or "")
              for r in details for c in (r.get("coupon_terms") or []) if c.get("couponCode")}
    # Terms differ by airline: each airline's come only from its own booking pages.
    own = st.terms_by_airline(details)
    auto_code = {(_route(r), r["airline"]): r.get("coupon_code") for r in search}

    out: list[dict[str, Any]] = []

    def where(r: dict[str, Any], market: str) -> dict[str, Any]:
        base = round(float(r.get("base_fare_bdt") or 0)) or None
        return {"market": market, "route": _route(r), "ota": "ShareTrip-B2C",
                "airline": r["airline"], "basis": BASE, "date": r.get("departure_date") or "",
                "fare": (_route(r), r.get("departure_date") or "", base)}

    for r in details:
        market = _market(r["flight_type"], r.get("origin"), r.get("destination"))
        out += _sharetrip_obs(where(r, market), float(r.get("base_pct") or 0),
                              auto_code.get((_route(r), r["airline"])),
                              r.get("coupon_terms") or [], r, titles)
    for r in search:
        market = _market(r["flight_type"], r.get("origin"), r.get("destination"))
        auto = float(r.get("discount_pct") or 0)
        terms = own.get(("DOM" if market == "DOM" else "INTL", r["airline"]))
        cell = (st.judge_cell(auto, float(r.get("base_fare_bdt") or 0), terms, gateways=gateways)
                if terms is not None else None)
        out += _sharetrip_obs(where(r, market), auto, r.get("coupon_code"), terms, cell, titles)
    return out


# --- GoZayaan -----------------------------------------------------------------------

def gozayaan(hars: Optional[list[str]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for h in hars or []:
        _rows, surcharge = g._recall("gozayaan_routed", h, ([], {}))
        rows = g._recall("gozayaan_fares", h, None)
        for r in rows if rows is not None else _rows:
            market = _market(r["flight_type"], r.get("origin"), r.get("destination"))
            date = r.get("departure_date") or ""
            common = r["eligibility_scope"] == "common"
            flat = r.get("discount_type") == "FLAT"
            out.append(_obs(
                market=market, route=_route(r), ota="Go Zayaan", airline=r["airline"],
                tier=COMMON if common else SPECIAL, code=r["coupon_code"],
                # open offers say how to pay (bKash, any VISA); bank offers name every card
                who=r["eligibility"] if common else (r.get("bank_cards") or r["eligibility"]),
                published=None if flat else r["discount_pct"], effective=r["realized_pct"],
                basis=TOTAL, fare=(_route(r), date, r.get("product_price")), date=date,
                cap=r.get("cap_bdt"), fee=(surcharge or {}).get(market),
                title=(f"Flat BDT {r['flat_bdt']:,} off. " if flat and r.get("flat_bdt") else "")
                + (r.get("name") or "")))
    return out


# --- FirstTrip B2C ------------------------------------------------------------------

def _ft_effective(rate: float, base: float, cap: Any) -> float:
    """Rate after its BDT cap at this fare (no cap or no fare: the rate itself)."""
    if not cap or base <= 0:
        return rate
    return round(min(math.floor(base * rate / 100), float(cap)) / base * 100, 2)


def ft_tier(code: str, airline: str, *, special_slot: bool = False) -> tuple[str, str]:
    """(tier, who) for a FirstTrip code, from its brand core (FTEBLDOM07 -> EBL):
    a known card/membership -> Special; a wallet -> Common; the airline's own or a bare
    route code (FTBSDOM) -> Common; anything else -> Unclear (or Special when FirstTrip
    itself put it in the special-coupon slot)."""
    ft = g.firsttrip
    core = ft._ft_coupon_core(code)
    if core in ft._FT_CARD_LABELS:
        return SPECIAL, ft._FT_CARD_LABELS[core]
    if core in ft._FT_WALLET_LABELS:
        return COMMON, f"{ft._FT_WALLET_LABELS[core]} payment"
    if special_slot:
        return SPECIAL, "Card holders (FirstTrip special offer)"
    if not core or core == str(airline).upper():
        return COMMON, "Anyone"
    return UNCLEAR, "Not a known card or wallet code: check the offer on FirstTrip"


def _ft_coupon_obs(where: dict[str, Any], c: dict[str, Any], w: dict[str, Any],
                   fees: dict[str, Any]) -> dict[str, Any]:
    """One catalogue coupon on one fare."""
    fo = g.firsttrip_offers
    tier, who = fo.audience(c)
    tier = COMMON if tier == "common" else SPECIAL
    rate = w["rate"]
    flat = rate["type"] == "F"
    return _obs(**where, tier=tier, code=c["code"], who=who,
                published=None if flat else rate["value"], effective=w["pct"], cap=rate["cap"],
                stacks=c["stacks_with_dynamic"],
                fee=fees.get("common" if tier == COMMON else "card"),
                title=(f"Flat BDT {rate['value']:,.0f} off. " if flat else "")
                + (c["description"] or c["title"]))


def firsttrip_b2c(rows_by_route: Optional[dict], fees: Optional[dict],
                  catalog: Optional[dict] = None) -> list[dict[str, Any]]:
    """Every FirstTrip option per fare: the dynamic discount, the search's auto-applied
    coupon, and (when the payment-page offer list was captured) every coupon that
    applies to the fare. Catalogue coupons that matched no fare seen are listed per
    airline from their rate table; telco perks without a rate are shown as such."""
    fees = fees or {}
    catalog = catalog or {}
    coupons = catalog.get("coupons") or []
    in_catalog = {c["code"] for c in coupons}
    used: set = set()
    out: list[dict[str, Any]] = []
    for rows in (rows_by_route or {}).values():
        for r in rows:
            airline = r["airline"]
            base = float(r.get("base_fare_bdt") or 0)
            date = str(r.get("departure") or "")[:10]
            where = dict(market=_market("", r.get("origin"), r.get("destination")),
                         route=_route(r), ota="Firsttrip-B2C", airline=airline, basis=BASE,
                         date=date, fare=(_route(r), date, r.get("flight_number"), round(base)))
            fee_for = {COMMON: fees.get("common"), SPECIAL: fees.get("card")}
            before = len(out)
            if (r.get("dynamic_rate") or 0) > 0:
                out.append(_obs(**where, tier=COMMON, code=r.get("dynamic_code") or NO_CODE,
                                who="Anyone (automatic)", published=r["dynamic_rate"],
                                effective=g.firsttrip_offers.dynamic_pct(r), fee=fees.get("common")))
            for c in coupons:
                w = g.firsttrip_offers.worth(c, airline, r.get("origin") or "", r.get("destination") or "", base)
                if w:
                    used.add((c["code"], airline, _route(r)))
                    out.append(_ft_coupon_obs(where, c, w, fees))
            code, rate = r.get("coupon_code") or "", float(r.get("headline_rate") or 0)
            if rate > 0 and code not in in_catalog:
                tier, who = ft_tier(code, airline)
                flat = r.get("coupon_type") == "F"
                # uncapped: the published rate (16), not FirstTrip's floored taka amount (15.97)
                capped = _ft_effective(rate, base, r.get("coupon_cap_bdt"))
                full = math.floor(base * rate / 100) if base > 0 else 0
                cap = float(r.get("coupon_cap_bdt") or 0)
                out.append(_obs(**where, tier=tier, code=code or NO_CODE, who=who,
                                published=None if flat else rate,
                                effective=rate if flat or not (cap and full > cap) else capped,
                                cap=r.get("coupon_cap_bdt"), fee=fee_for.get(tier),
                                title=f"Flat BDT {r['coupon_flat_bdt']:,.0f} off" if flat and r.get("coupon_flat_bdt") else ""))
            slot, slot_rate = r.get("special_code") or "", float(r.get("special_rate") or 0)
            if slot_rate > 0 and slot != code and slot not in in_catalog:
                tier, who = ft_tier(slot, airline, special_slot=True)
                out.append(_obs(**where, tier=tier, code=slot or NO_CODE, who=who,
                                published=slot_rate, effective=slot_rate, fee=fee_for.get(tier)))
            if len(out) == before:     # e.g. VQ: on sale, nothing off; keep the airline visible
                out.append(_obs(**where, tier=NO_DISCOUNT, code="(none)", who="No discount on this fare",
                                published=0.0, effective=None,
                                title="FirstTrip showed this fare with no discount or coupon."))
    out += _ft_telco_obs(rows_by_route, catalog, fees)
    out += _ft_rate_table_obs(coupons, used, fees)
    out += _ft_perk_obs(catalog.get("perks") or {}, catalog.get("perk_offers") or [])
    return out


def _ft_telco_obs(rows_by_route: Optional[dict], catalog: dict, fees: dict[str, Any]) -> list[dict[str, Any]]:
    """Verified telco offers (FTGPSTAR...) on each fare of the airline they were verified for;
    an offer whose airline had no fare in this run is listed once from its own capture."""
    fo = g.firsttrip_offers
    out, placed = [], set()
    for rows in (rows_by_route or {}).values():
        for r in rows:
            base = float(r.get("base_fare_bdt") or 0)
            date = str(r.get("departure") or "")[:10]
            for p in fo.perks_for(catalog, _market("", r.get("origin"), r.get("destination")), r["airline"]):
                w = fo.perk_worth(p, base)
                if not w:
                    continue
                placed.add((p["code"], p["market"], p["airline"]))
                out.append(_telco_obs(p, market=_market("", r.get("origin"), r.get("destination")),
                                      route=_route(r), effective=w["pct"], date=date,
                                      fare=(_route(r), date, r.get("flight_number"), round(base)),
                                      fee=fees.get("card")))
    for p in catalog.get("perk_offers") or []:
        if (p["code"], p["market"], p["airline"]) not in placed:
            route = f"{p['origin']}-{p['destination']}" if p["origin"] and p["destination"] else ALL_ROUTES
            out.append(_telco_obs(p, market=p["market"], route=route, effective=None,
                                  fee=fees.get("card")))
    return out


def _telco_obs(p: dict[str, Any], *, market: str, route: str, effective: Any, fee: Any,
               date: str = "", fare: Any = None) -> dict[str, Any]:
    flat = p["type"] == "F"
    return _obs(market=market, route=route, ota="Firsttrip-B2C", airline=p["airline"],
                tier=SPECIAL, code=p["code"], who=f"{p['operator']} customers (number verified by OTP)",
                published=None if flat else p["value"], effective=effective, basis=BASE,
                cap=p["cap"], stacks=False, fee=fee, date=date, fare=fare,
                title=(f"Flat BDT {p['value']:,.0f} off. " if flat else "")
                + (p["description"] or "Telco perk") + ". Replaces the automatic discount.")


def _ft_rate_table_obs(coupons: list[dict[str, Any]], used: set,
                       fees: dict[str, Any]) -> list[dict[str, Any]]:
    """Coupon rates from the offer list for airline/routes with no fare in this run: the
    coupon's rate table still says what each gets (one row per table entry)."""
    out = []
    fo = g.firsttrip_offers
    for c in coupons:
        tier, who = fo.audience(c)
        for d in c["config"]:
            route = f"{d['origin']}-{d['destination']}" if d["origin"] and d["destination"] else ALL_ROUTES
            seen_airline = any(u[0] == c["code"] and u[1] == d["airline"] for u in used)
            if not d["airline"] or (c["code"], d["airline"], route) in used \
                    or (route == ALL_ROUTES and seen_airline):
                continue
            flat = d["type"] == "F"
            out.append(_obs(market=c["market"], route=route, ota="Firsttrip-B2C", airline=d["airline"],
                            tier=NOT_SEARCHED, code=c["code"],
                            who=who + ("" if tier == "common" else " (card offer)"),
                            published=None if flat else d["value"], effective=None, basis=BASE,
                            cap=d["cap"] or c["cap"], stacks=c["stacks_with_dynamic"],
                            fee=fees.get("common" if tier == "common" else "card"),
                            title=("Route not searched in this run: this rate comes from "
                                   "FirstTrip's coupon table. The route's automatic discount and "
                                   "its other offers show once it is searched (live or HAR). "
                                   + (f"Flat BDT {d['value']:,.0f} off. " if flat else "")
                                   + (c["description"] or ""))))
    return out


def _ft_perk_obs(perks: dict[str, list[str]], offers: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """FirstTrip's telco perks whose rate wasn't captured: the partners are listed for the
    fare, but a rate only appears after the customer picks an operator and verifies the
    number (one capture per operator)."""
    out = []
    for market, names in perks.items():
        have = {p["operator"] for p in offers if p["market"] == market}
        missing = [n for n in names if n not in have]
        if missing:
            out.append(_obs(market=market, route=ALL_ROUTES, ota="Firsttrip-B2C", airline="All",
                            tier=NOT_CAPTURED, code="(telco perk)", who=f"{' / '.join(missing)} customers",
                            published=None, effective=None, basis=BASE,
                            title="FirstTrip offers these telco perks, but the rate shows only after "
                                  "choosing the operator and verifying a phone number of that "
                                  "operator on the payment page. Capture that step once per operator."))
    return out


# --- merge + entry point ------------------------------------------------------------

def _fmt_range(lo: Optional[float], hi: Optional[float]) -> str:
    if hi is None:
        return ""
    return g._fmt(hi) if lo is None or round(lo, 2) == round(hi, 2) else f"{g._fmt(lo)}-{g._fmt(hi)}"


def _fmt_dates(dates: set) -> str:
    ds = sorted(d for d in dates if d)
    if len(ds) <= 3:
        return ", ".join(ds)
    return f"{ds[0]} to {ds[-1]} ({len(ds)} dates)"


def merge(observations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One entry per (market, route, OTA, airline, tier, code): the effective rate as a
    lowest-highest range over the DISTINCT fares seen, sorted by tier, then best first."""
    by_key: dict[tuple, dict[str, Any]] = {}
    for o in observations:
        key = (o["market"], o["route"], o["ota"], o["airline"], o["tier"], o["code"])
        e = by_key.setdefault(key, {**o, "eff_lo": None, "eff_hi": None, "fares": set(),
                                    "dates": set(), "unkeyed": 0})
        eff = o["effective_pct"]
        if eff is not None:
            e["eff_lo"] = eff if e["eff_lo"] is None else min(e["eff_lo"], eff)
            e["eff_hi"] = eff if e["eff_hi"] is None else max(e["eff_hi"], eff)
        if o.get("fare") is None:
            e["unkeyed"] += 1
        else:
            e["fares"].add(o["fare"])
        if o.get("date"):
            e["dates"].add(o["date"])
        for k in ("who", "title", "cap_bdt", "fee_pct", "published_pct", "stacks"):
            if e.get(k) in (None, "") and o.get(k) not in (None, ""):
                e[k] = o[k]
    entries = []
    for e in by_key.values():
        fares, dates = e.pop("fares"), e.pop("dates")
        e["seen"] = len(fares) + e.pop("unkeyed")
        e["dates"] = sorted(d for d in dates if d)
        e["travel_dates"] = _fmt_dates(dates)
        for k in ("effective_pct", "fare", "date"):
            e.pop(k, None)
        e["effective"] = _fmt_range(e["eff_lo"], e["eff_hi"])
        e["stacks"] = "" if e["stacks"] is None else ("Yes" if e["stacks"] else "No")
        entries.append(e)
    ota_order = {lab: i for i, (lab, _k) in enumerate(g.ROW_ORDER)}
    entries.sort(key=lambda e: (e["market"] != "DOM", e["route"] in ("", ALL_ROUTES), e["route"],
                                ota_order.get(e["ota"], 99), e["airline"],
                                TIER_ORDER.get(e["tier"], 9), -(e["eff_hi"] or 0), e["code"]))
    return entries


_WIDTHS = {"market": 13, "route": 10, "travel_dates": 14, "ota": 15, "airline": 8, "tier": 12,
           "code": 18, "who": 30, "published_pct": 11, "effective": 12, "basis": 13,
           "cap_bdt": 10, "stacks": 11, "fee_pct": 10, "seen": 8, "title": 60}
_NOTE = ("Every promo code each OTA offered, not just the best one shown in the grids. "
         "Common = anyone gets it; Special = needs that card, bank or membership; "
         "Not captured = ShareTrip's coupons differ by airline and this airline's booking "
         "page wasn't captured; Not searched = a FirstTrip coupon-table rate for a route "
         "this run didn't search (its automatic discount and other offers are unknown "
         "until it is searched); No discount = a 0% payment option (EMI, net banking). "
         "Effective % = what the code was worth at the fares seen (caps applied; a range "
         "when fares differed), including the automatic discount when the code goes on "
         "top of it. '% of' says what the % is taken from: ShareTrip and FirstTrip use the "
         "base fare, GoZayaan the booking total, so compare like with like. "
         "Use the filter arrows in the header row.")
_TIER_FILL = {COMMON: "C6EFCE", SPECIAL: "DDEBF7", UNCLEAR: "FFF2CC",
              NOT_CAPTURED: "FFC7CE", NOT_SEARCHED: "EDEDED", NO_DISCOUNT: "EDEDED"}


def _pct(v: Any) -> Any:
    """'12.5' -> 0.125 (shown as a %); a range like '7.41-18' stays text."""
    if v in (None, ""):
        return None
    try:
        return float(v) / 100.0
    except (TypeError, ValueError):
        return f"{v}%"


def write_sheet(ws, report: dict[str, Any]) -> None:
    """The (all codes) sheet: one filterable row per code."""
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter
    st = g._detail_styles()
    ncol = len(COLUMNS)
    tier_col = [k for k, _h in COLUMNS].index("tier") + 1
    t = ws.cell(1, 1, f"{report['report_date']} / {report['report_time']}hrs — all promo codes")
    t.font = st["title"]
    ws.cell(2, 1, _NOTE).font = st["note"]
    for ci, (key, head) in enumerate(COLUMNS, start=1):
        c = ws.cell(3, ci, head)
        c.font, c.fill, c.alignment, c.border = st["head"], st["hdr_fill"], st["center"], st["border"]
        ws.column_dimensions[get_column_letter(ci)].width = _WIDTHS[key]
    r = 4
    for e in report.get("all_codes") or []:
        vals = {**e, "market": "Domestic" if e["market"] == "DOM" else "International",
                "published_pct": _pct(e.get("published_pct")), "effective": _pct(e.get("effective")),
                "fee_pct": _pct(e.get("fee_pct"))}
        for ci, (key, _h) in enumerate(COLUMNS, start=1):
            v = vals.get(key)
            cell = ws.cell(r, ci, v if v != "" else None)
            cell.font = st["label"] if key == "code" else st["data"]
            cell.alignment = st["left"] if key in ("code", "who", "title", "ota", "travel_dates") \
                else st["center"]
            cell.border = st["border"]
            if isinstance(v, float) and key in ("published_pct", "effective", "fee_pct"):
                cell.number_format = "0.##%"
            elif key == "cap_bdt" and v:
                cell.number_format = "#,##0"
        if e["tier"] in _TIER_FILL:
            ws.cell(r, tier_col).fill = PatternFill("solid", fgColor=_TIER_FILL[e["tier"]])
        r += 1
    ws.auto_filter.ref = f"A3:{get_column_letter(ncol)}{max(r - 1, 3)}"
    ws.freeze_panes = "A4"
    ws.print_title_rows = "3:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def collect(*, sharetrip_hars=None, gozayaan_hars=None, b2c_rows_by_route=None,
            b2c_fees=None) -> list[dict[str, Any]]:
    """Every code from every coupon channel. Must run inside build_report (the parse
    memo is live). One channel failing only drops that channel from the list."""
    observations: list[dict[str, Any]] = []
    steps = [("ShareTrip", lambda: sharetrip(sharetrip_hars)),
             ("GoZayaan", lambda: gozayaan(gozayaan_hars)),
             ("FirstTrip B2C", lambda: firsttrip_b2c(b2c_rows_by_route, b2c_fees,
                                                     g._recall("ft_catalog", "all")))]
    for name, step in steps:
        try:
            observations += step()
        except Exception as exc:  # noqa: BLE001 — surfaced in the run log
            print(f"  ! all-codes list: {name} skipped: {exc}")
    return merge(observations)
