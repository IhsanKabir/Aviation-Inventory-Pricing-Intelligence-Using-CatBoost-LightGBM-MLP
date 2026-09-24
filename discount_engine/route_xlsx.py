"""Excel sheets for the per-route view: overview, detail, and flat data.

  * <day> routes       - the answer at a glance: per route and airline, the best
                         customer discount (B2C) and the best agent commission (B2B),
                         each with its OTA; then which OTAs covered which route.
  * <day> route detail - one block per route (OTA rows x airline columns), blank OTA
                         rows dropped, coupon tiers on their own lines, sized so no
                         text is clipped, a page break before every block.
  * <day> route data   - one row per route x OTA x airline, filterable.

Why B2C and B2B are shown apart in the overview: an agent commission (BDFare) and a
customer discount (ShareTrip) are different money; the grid already ranks them in
separate groups, so the overview must not crown one "best" across both.

Routes come from report["by_route"] (already highlighted), optionally ordered by
with_preferred(): preferred routes first, the rest under "Other routes found".
"""
from __future__ import annotations

import math
import re
from typing import Any, Optional

from .highlight import BEST_SHORT, _split_top_level, parse_cell_tiers

# Palette shared with grid.py's daily sheet (import would be circular at call time).
HDR_BG, HDR_TX = "1F4E79", "FFFFFF"
BEST_BG, BEST_TX = "1F2937", "FFFFFF"
HL_GREEN, HL_GREEN_TX = "C6EFCE", "006100"
HL_BLUE, HL_BLUE_TX = "DDEBF7", "1F4E79"
GREY_BG, GREY_TX = "F2F2F2", "7F7F7F"
PART_BG = "FFF2CC"                       # partial coverage
LEGEND = "Green = best  ·  Blue = 2nd  (ranked within B2B and within B2C, net of fees)"
LINE_PT = 15                             # row height per text line (Calibri 11)

_PCT = re.compile(r"(?<![A-Za-z\d.])(\d+(?:\.\d+)?)(?![\d.%A-Za-z])")


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="BFBFBF")
    cal = "Calibri"
    return {
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "hdr_fill": PatternFill("solid", fgColor=HDR_BG),
        "best_fill": PatternFill("solid", fgColor=BEST_BG),
        "green": PatternFill("solid", fgColor=HL_GREEN),
        "blue": PatternFill("solid", fgColor=HL_BLUE),
        "grey": PatternFill("solid", fgColor=GREY_BG),
        "part": PatternFill("solid", fgColor=PART_BG),
        "title": Font(name=cal, bold=True, size=13),
        "h2": Font(name=cal, bold=True, size=12, color=HDR_BG),
        "link": Font(name=cal, bold=True, size=11, color="0563C1", underline="single"),
        "head": Font(name=cal, bold=True, size=11, color=HDR_TX),
        "label": Font(name=cal, bold=True, size=11),
        "data": Font(name=cal, size=11),
        "note": Font(name=cal, italic=True, size=10, color=GREY_TX),
        "green_tx": Font(name=cal, bold=True, size=11, color=HL_GREEN_TX),
        "blue_tx": Font(name=cal, size=11, color=HL_BLUE_TX),
        "grey_tx": Font(name=cal, size=11, color=GREY_TX),
        "best": Font(name=cal, bold=True, size=11, color=BEST_TX),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "left": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "left_top": Alignment(horizontal="left", vertical="top", wrap_text=True),
    }


# ----------------------------------------------------------------------------- helpers

def _has_data(row: dict[str, Any]) -> bool:
    return any(str(v or "").strip() for v in (row.get("cells") or {}).values())


def visible_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Rows with data; a B2B/B2C separator only when both sides kept a row."""
    out: list[dict[str, Any]] = []
    for row in rows:
        if row.get("kind") == "sep":
            if out and out[-1].get("kind") != "sep":
                out.append(row)
        elif _has_data(row):
            out.append(row)
    while out and out[-1].get("kind") == "sep":
        out.pop()
    return out


def cell_lines(raw: Any) -> list[str]:
    """'9.5(Bkash, 2% fee), 18 (Stellar, 2% fee)' -> ['9.5% (Bkash, 2% fee)', '18% (Stellar, 2% fee)']."""
    text = str(raw or "").strip()
    if not text:
        return []
    parts = _split_top_level(text) or [text]
    out = []
    for part in parts:
        part = _PCT.sub(r"\1%", part, count=1)
        out.append(re.sub(r"%\(", "% (", part))
    return out


def _is_plain_number(raw: str) -> Optional[float]:
    try:
        return float(raw)
    except ValueError:
        return None


def _write_rate(cell, raw: str, st) -> list[str]:
    """Plain rate -> number in % format; tiered text -> one tier per line."""
    raw = str(raw or "").strip()
    num = _is_plain_number(raw)
    if num is not None:
        cell.value, cell.number_format = num / 100.0, "0.00%"
        return [f"{num:g}%"]
    lines = cell_lines(raw)
    cell.value = "\n".join(lines) or None
    return lines


def _height(lines_by_col: dict[int, list[str]], widths: dict[int, float], minimum: int = 18) -> float:
    """Row height that shows every wrapped line: per cell, each line wraps at its
    column width; the tallest cell sets the row."""
    tallest = 1
    for col, lines in lines_by_col.items():
        chars = max(4, int(widths.get(col, 12) * 1.1) - 1)
        tallest = max(tallest, sum(max(1, math.ceil(len(l) / chars)) for l in lines) or 1)
    return max(minimum, tallest * LINE_PT + 3)


def _groups(block: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    out: dict[str, list[dict[str, Any]]] = {"b2b": [], "b2c": []}
    for row in block.get("rows", []):
        if row.get("kind") in out and _has_data(row):
            out[row["kind"]].append(row)
    return out


def best_in_group(rows: list[dict[str, Any]], airline: str) -> Optional[dict[str, Any]]:
    """Best NET common rate for one airline among one group's rows (B2B or B2C)."""
    best = None
    for row in rows:
        raw = str((row.get("cells") or {}).get(airline) or "").strip()
        tiers = parse_cell_tiers(raw)
        if not tiers:
            continue
        cand = {"net": tiers[0]["net"], "pct": tiers[0]["pct"], "label": row["label"],
                "estimate": raw.startswith("~")}
        if best is None or cand["net"] > best["net"]:
            best = cand
    return best


def _market_name(market: str) -> str:
    return "Domestic" if market == "DOM" else "International"


def _sections(report: dict[str, Any]) -> list[tuple[Optional[str], list[dict[str, Any]]]]:
    """[(heading or None, blocks)] - preferred first when the user has any."""
    blocks = report.get("by_route") or []
    preferred = [b for b in blocks if b.get("preferred")]
    missing = report.get("preferred_missing") or []
    if not preferred and not missing:
        return [(None, blocks)]
    others = [b for b in blocks if not b.get("preferred")]
    out = [("★ Your preferred routes", preferred)]
    if others:
        out.append(("Other routes found in this run", others))
    return out


def _stamp(report: dict[str, Any]) -> str:
    return f"{report.get('report_date', '')} · {report.get('report_time', '')}hrs"


#: Wider than this and "fit to one page wide" shrinks the text past reading size.
FIT_MAX_COLUMNS = 12


def _print_setup(ws, ncols: int) -> None:
    """Landscape, fitted to the page width - one page wide for narrow sheets, two
    for wide ones (up to 26 international airlines), with column A repeated on the
    second. One page wide shrank wide sheets past reading size; a fixed scale split
    each wide block across three pages (28 pages for one day's detail)."""
    ws.page_setup.orientation = "landscape"
    ws.print_title_cols = "A:A"
    ws.page_setup.fitToWidth = 1 if ncols <= FIT_MAX_COLUMNS else 2
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ----------------------------------------------------------------------------- detail

def render_detail(ws, report: dict[str, Any], overview_title: str) -> dict[str, int]:
    """One block per route. Returns {route: first row} for the overview's links."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break

    st = _styles()
    blocks = [b for _h, bs in _sections(report) for b in bs]
    used = {lab for b in blocks for lab in b.get("coverage", {}).get("with_data", [])}

    # Column widths across ALL blocks (they share the sheet's columns): wide enough
    # for the longest line in that column, capped so wrapping takes over.
    longest: dict[int, int] = {}
    for b in blocks:
        for ci, airline in enumerate(b["columns"], start=2):
            longest[ci] = max(longest.get(ci, 6), len(airline) + 2)
            for row in visible_rows(b["rows"]):
                for line in cell_lines((row.get("cells") or {}).get(airline)):
                    longest[ci] = max(longest[ci], len(line))
    widths = {ci: min(30, max(11, n * 0.95 + 2)) for ci, n in longest.items()}
    widths[1] = 18
    for ci, w in widths.items():
        ws.column_dimensions[get_column_letter(ci)].width = w
    span = max([1] + list(widths))

    ws.cell(1, 1, f"Discounts by route — detail · {_stamp(report)}").font = st["title"]
    ws.cell(2, 1, LEGEND).font = st["note"]
    back = ws.cell(3, 1, "↑ Back to the routes overview")
    back.hyperlink, back.font = f"#'{overview_title}'!A1", st["note"]
    r, anchors = 5, {}
    first_block = True
    for heading, section in _sections(report):
        break_done = False
        if heading:
            if not first_block:                             # heading starts a new page,
                ws.row_breaks.append(Break(id=r - 1))       # it never trails the last one
                break_done = True
            ws.cell(r, 1, heading).font = st["h2"]
            r += 1
        if heading and heading.startswith("★") and report.get("preferred_missing"):
            ws.cell(r, 1, "Not captured yet: " + ", ".join(report["preferred_missing"])).font = st["note"]
            r += 1
        for b in section:
            if not first_block and not break_done:
                ws.row_breaks.append(Break(id=r - 1))       # never split a block
            first_block, break_done = False, False
            anchors[b["route"]] = r
            cols = b["columns"]
            ncol = 1 + len(cols)
            title = f"{b['route']} · {_market_name(b['market'])}" + ("  ★" if b.get("preferred") else "")
            t = ws.cell(r, 1, title)
            t.font, t.alignment = st["title"], st["left"]
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(ncol, 2))
            ws.row_dimensions[r].height = 22
            r += 1
            missing = [lab for lab in b.get("coverage", {}).get("without", []) if lab in used]
            if missing:
                n = ws.cell(r, 1, "No data on this route: " + ", ".join(missing))
                n.font, n.alignment = st["note"], st["left"]
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(ncol, 2))
                r += 1
            for ci, head in enumerate(["OTA", *cols], start=1):
                c = ws.cell(r, ci, head)
                c.font, c.fill, c.alignment, c.border = st["head"], st["hdr_fill"], st["center"], st["border"]
            r += 1
            for row in visible_rows(b["rows"]):
                if row.get("kind") == "sep":
                    r += 1
                    continue
                lab = ws.cell(r, 1, row["label"])
                lab.font, lab.alignment, lab.border = st["label"], st["left"], st["border"]
                lines_by_col = {}
                for ci, airline in enumerate(cols, start=2):
                    cell = ws.cell(r, ci)
                    cell.font, cell.alignment, cell.border = st["data"], st["center"], st["border"]
                    lines_by_col[ci] = _write_rate(cell, (row.get("cells") or {}).get(airline), st)
                    flag = (row.get("highlights") or {}).get(airline)
                    if flag == "highest":
                        cell.fill, cell.font = st["green"], st["green_tx"]
                    elif flag == "second":
                        cell.fill, cell.font = st["blue"], st["blue_tx"]
                ws.row_dimensions[r].height = _height(lines_by_col, widths)
                r += 1
            # Best row (same meaning as the exe and the daily sheet)
            bl = ws.cell(r, 1, "Best (net)")
            bl.font, bl.fill, bl.alignment, bl.border = st["best"], st["best_fill"], st["left"], st["border"]
            lines_by_col = {}
            for ci, airline in enumerate(cols, start=2):
                b_ = (b.get("best") or {}).get(airline)
                lines = []
                if b_:
                    lines.append(b_["universal"]["display"])
                    g = b_.get("gated")
                    if g and g["net"] > b_["universal"]["net"]:
                        lines.append("card: " + g["display"])
                c = ws.cell(r, ci, "\n".join(lines) or None)
                c.font, c.fill, c.alignment, c.border = st["best"], st["best_fill"], st["center"], st["border"]
                lines_by_col[ci] = lines
            ws.row_dimensions[r].height = _height(lines_by_col, widths)
            r += 3
    ws.freeze_panes = "B1"
    _print_setup(ws, span)
    return anchors


# ----------------------------------------------------------------------------- overview

def _best_lines(block: dict[str, Any], airline: str) -> tuple[list[str], Optional[float]]:
    """Two labelled lines (B2C, B2B) and the highest net among them (for colour)."""
    groups = _groups(block)
    lines, nets = [], []
    for kind, tag in (("b2c", "B2C"), ("b2b", "B2B")):
        best = best_in_group(groups[kind], airline)
        if not best:
            continue
        est = "~" if best["estimate"] else ""
        short = BEST_SHORT.get(best["label"], best["label"])
        if best["net"] <= 0:
            lines.append(f"{tag} no discount · {short}")
        else:
            lines.append(f"{tag} {est}{best['pct']:g}% · {short}")
        nets.append(best["net"])
    return lines, (max(nets) if nets else None)


def render_overview(ws, report: dict[str, Any], detail_title: str,
                    anchors: dict[str, int]) -> None:
    from openpyxl.utils import get_column_letter

    st = _styles()
    ws.cell(1, 1, f"Discounts by route — overview · {_stamp(report)}").font = st["title"]
    ws.cell(2, 1, "Each cell: the best customer discount (B2C) and the best agent commission "
                  "(B2B) for that airline on that route, with the OTA. They are different money, "
                  "so they are never ranked against each other. ~ = estimated base. "
                  "Click a route for its full detail.").font = st["note"]
    blocks = [b for _h, bs in _sections(report) for b in bs]
    r = 4
    widths: dict[int, float] = {1: 14}
    for market in ("DOM", "INTL"):
        mblocks = [b for b in blocks if b["market"] == market]
        missing = [m for m in (report.get("preferred_missing") or [])
                   if _route_market(m) == market]
        if not mblocks and not missing:
            continue
        airlines = _market_airlines(mblocks)
        ws.cell(r, 1, f"{_market_name(market)} — best discount by route").font = st["h2"]
        r += 1
        for ci, head in enumerate(["Route", *airlines], start=1):
            c = ws.cell(r, ci, head)
            c.font, c.fill, c.alignment, c.border = st["head"], st["hdr_fill"], st["center"], st["border"]
            widths[ci] = max(widths.get(ci, 0), 20 if ci > 1 else 14)
        r += 1
        for heading, section in _sections(report):
            section = [b for b in section if b["market"] == market]
            sec_missing = missing if heading and heading.startswith("★") else []
            if heading and (section or sec_missing):
                h = ws.cell(r, 1, heading)
                h.font = st["label"]
                r += 1
            for b in section:
                rc = ws.cell(r, 1, b["route"] + (" ★" if b.get("preferred") else ""))
                rc.font, rc.alignment, rc.border = st["link"], st["left"], st["border"]
                rc.hyperlink = f"#'{detail_title}'!A{anchors.get(b['route'], 1)}"
                lines_by_col = {}
                for ci, airline in enumerate(airlines, start=2):
                    cell = ws.cell(r, ci)
                    cell.alignment, cell.border = st["center"], st["border"]
                    if airline not in b["columns"]:
                        cell.font = st["data"]
                        continue
                    lines, top = _best_lines(b, airline)
                    cell.value = "\n".join(lines) or None
                    lines_by_col[ci] = lines
                    if top is None:
                        cell.font = st["data"]
                    elif top > 0:
                        cell.fill, cell.font = st["green"], st["green_tx"]
                    else:
                        cell.fill, cell.font = st["grey"], st["grey_tx"]
                ws.row_dimensions[r].height = _height(lines_by_col, widths)
                r += 1
            for route in sec_missing:
                rc = ws.cell(r, 1, route + " ★")
                rc.font, rc.border = st["grey_tx"], st["border"]
                n = ws.cell(r, 2, "not captured yet")
                n.font = st["note"]
                r += 1
        r += 2

    # Coverage: how many of the route's airlines each OTA priced.
    labels = [lab for lab in BEST_SHORT if any(
        lab in b.get("coverage", {}).get("with_data", []) for b in blocks)]
    if blocks and labels:
        ws.cell(r, 1, "Coverage — airlines priced by each OTA on each route").font = st["h2"]
        r += 1
        for ci, head in enumerate(["Route", *labels], start=1):
            c = ws.cell(r, ci, head)
            c.font, c.fill, c.alignment, c.border = st["head"], st["hdr_fill"], st["center"], st["border"]
            widths[ci] = max(widths.get(ci, 0), 14)
        ws.row_dimensions[r].height = 30
        r += 1
        for b in blocks:
            rc = ws.cell(r, 1, b["route"] + (" ★" if b.get("preferred") else ""))
            rc.font, rc.border = st["label"], st["border"]
            total = len(b["columns"])
            rows = {row["label"]: row for row in b["rows"] if row.get("kind") != "sep"}
            for ci, lab in enumerate(labels, start=2):
                n = sum(1 for v in (rows.get(lab, {}).get("cells") or {}).values() if str(v or "").strip())
                c = ws.cell(r, ci, f"{n}/{total}" if n else "–")
                c.alignment, c.border = st["center"], st["border"]
                if n == total:
                    c.fill, c.font = st["green"], st["green_tx"]
                elif n:
                    c.fill, c.font = st["part"], st["data"]
                else:
                    c.fill, c.font = st["grey"], st["grey_tx"]
            r += 1

    for ci, w in widths.items():
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "B1"
    _print_setup(ws, max(widths))


def _route_market(route: str) -> str:
    from .grid import _route_type
    o, _, d = str(route).partition("-")
    return _route_type(o, d)


def _market_airlines(blocks: list[dict[str, Any]]) -> list[str]:
    from .grid import DOM_COLUMNS, INTL_COLUMNS
    present = {a for b in blocks for a in b["columns"]}
    base = DOM_COLUMNS if blocks and blocks[0]["market"] == "DOM" else INTL_COLUMNS
    return [a for a in base if a in present] + sorted(present - set(base))


# ----------------------------------------------------------------------------- data

DATA_HEADERS = ["Route", "Market", "Preferred", "OTA", "Type", "Airline", "Rate %",
                "Offer / method", "Fee %", "Net %", "Card %", "Card / loyalty", "Capped",
                "Estimated", "Best in group"]


def render_data(ws, report: dict[str, Any]) -> None:
    from openpyxl.utils import get_column_letter

    st = _styles()
    for ci, h in enumerate(DATA_HEADERS, start=1):
        c = ws.cell(1, ci, h)
        c.font, c.fill, c.alignment, c.border = st["head"], st["hdr_fill"], st["center"], st["border"]
    r = 2
    for _h, section in _sections(report):
        for b in section:
            for row in b.get("rows", []):
                if row.get("kind") == "sep":
                    continue
                for airline in b["columns"]:
                    raw = str((row.get("cells") or {}).get(airline) or "").strip()
                    tiers = parse_cell_tiers(raw)
                    if not tiers:
                        continue
                    common = tiers[0]
                    card = max(tiers[1:], key=lambda t: t["net"]) if len(tiers) > 1 else None
                    flag = (row.get("highlights") or {}).get(airline)
                    vals = [b["route"], _market_name(b["market"]), "yes" if b.get("preferred") else "",
                            row["label"], "B2B" if row.get("kind") == "b2b" else "B2C", airline,
                            common["pct"] / 100, common["label"] or "", (common["fee_pct"] / 100) or None,
                            common["net"] / 100, (card["pct"] / 100) if card else None,
                            card["label"] if card else "", "yes" if card and card["capped"] else "",
                            "yes" if raw.startswith("~") else "", "yes" if flag == "highest" else ""]
                    for ci, v in enumerate(vals, start=1):
                        c = ws.cell(r, ci, v if v != "" else None)
                        c.font, c.border = st["data"], st["border"]
                        if ci in (7, 9, 10, 11) and isinstance(v, float):
                            c.number_format = "0.00%"
                    r += 1
    widths = [10, 13, 10, 15, 7, 8, 9, 18, 8, 9, 9, 20, 8, 10, 12]
    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "A2"
    if r > 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(DATA_HEADERS))}{r - 1}"


def add_route_sheets(wb, report: dict[str, Any], day: str) -> None:
    """Overview, detail and data sheets, in that order, when the report has routes."""
    if not report.get("by_route"):
        return
    overview_title, detail_title = f"{day} routes", f"{day} route detail"
    overview = wb.create_sheet(title=overview_title)
    detail = wb.create_sheet(title=detail_title)
    anchors = render_detail(detail, report, overview_title)
    render_overview(overview, report, detail_title, anchors)
    render_data(wb.create_sheet(title=f"{day} route data"), report)
