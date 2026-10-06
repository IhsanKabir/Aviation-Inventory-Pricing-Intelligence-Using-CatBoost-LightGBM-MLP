"""All codes by TYPE: one row per route x OTA x airline, discount options as columns.

all_codes.merge() gives one entry per code. This view regroups them so the options
line up across OTAs:

  Automatic | bKash | Nagad & other wallets | Telco | Any card / payment | Bank & card offers

Each cell lists that type's codes on the row, best first ("12% FT-bKASH"). Two summary
columns give the best rate anyone can get and the best with a card or telco code.
0% payment options (EMI, net banking) are left out here; the detail sheet keeps them.
"""
from __future__ import annotations

from typing import Any, Optional

from . import all_codes as ac

AUTO, BKASH, WALLETS, TELCO, ANY_CARD, BANK, OTHER = (
    "Automatic", "bKash", "Nagad & other wallets", "Telco", "Any card / payment",
    "Bank & card offers", "Other")
KINDS = [AUTO, BKASH, WALLETS, TELCO, ANY_CARD, BANK, OTHER]
_WALLET_WORDS = ("nagad", "upay", "rocket", "tap", "wallet", "cellfin")


def kind_of(e: dict[str, Any]) -> Optional[str]:
    """Which column an all-codes entry belongs in (None = not shown in this view)."""
    tier, who, code = e["tier"], str(e.get("who") or "").lower(), str(e.get("code") or "")
    if tier == ac.NO_DISCOUNT:
        return None
    if code == "(telco perk)" or "verified" in who or who.startswith("anyone entering"):
        return TELCO
    if tier == ac.NOT_CAPTURED:
        return None                                   # a row note, not an option
    if who.startswith("anyone (automatic)"):
        return AUTO
    open_to_all = tier == ac.COMMON or (tier == ac.NOT_SEARCHED and "(card offer)" not in who)
    if open_to_all:
        if "any card" in who:                         # broadest audience wins (FTINT26: wallets or any card)
            return ANY_CARD
        if "bkash" in who:
            return BKASH
        if any(w in who for w in _WALLET_WORDS):
            return WALLETS
        return ANY_CARD
    return OTHER if tier == ac.UNCLEAR else BANK


def _value(e: dict[str, Any]) -> Optional[float]:
    return e.get("eff_hi") if e.get("eff_hi") is not None else e.get("published_pct")


def _who_short(e: dict[str, Any], limit: int = 34) -> str:
    who = str(e.get("who") or "")
    for prefix in ("Pay with ", "Anyone entering the "):
        if who.startswith(prefix):
            who = who[len(prefix):]
    who = who.replace(" code (not verified)", "").replace(" customers (number verified by OTP)", "")
    return who if len(who) <= limit else who[: limit - 1].rstrip(" ;,/") + "…"


def item_text(e: dict[str, Any], kind: str) -> str:
    """'12% FT-bKASH' / '13% RAMEX0126PBB13O (City Bank AMEX)' / 'rate not captured (Robi)'."""
    if e["code"] == "(telco perk)":
        return f"rate not captured ({_who_short(e).replace(' customers', '')})"
    eff = e.get("effective") or (ac.g._fmt(e["published_pct"]) if e.get("published_pct") is not None else "")
    text = f"{eff}% {e['code']}" if eff else e["code"]
    if kind in (TELCO, BANK, OTHER):
        text += f" ({_who_short(e)})"
    return text


def rows(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Regroup all-codes entries into one row per (market, route, OTA, airline, coverage)."""
    by_key: dict[tuple, dict[str, Any]] = {}
    for e in entries:
        kind = kind_of(e)
        note = ""
        if e["tier"] == ac.NOT_CAPTURED and kind is None:
            note = "Coupons not captured (booking page)"
        elif e["tier"] == ac.NO_DISCOUNT and e["route"] != ac.ALL_ROUTES:
            note = "No discount offered"              # an airline searched with nothing off: keep its row
        if kind is None and not note:
            continue
        status = "Not searched" if e["tier"] == ac.NOT_SEARCHED else ""
        key = (e["market"], e["route"], e["ota"], e["airline"], status)
        row = by_key.setdefault(key, {
            "market": e["market"], "route": e["route"], "ota": e["ota"], "airline": e["airline"],
            "status": status, "dates": set(), "seen": 0, "cells": {k: [] for k in KINDS},
            "best_anyone": None, "best_any": None, "notes": [], "basis": e.get("basis") or ""})
        row["dates"].update(e.get("dates") or [])
        row["seen"] = max(row["seen"], e.get("seen") or 0)
        if note:
            row["notes"].append(note)
            continue
        row["cells"][kind].append(e)
        v = _value(e)
        if v is None:
            continue
        if e["tier"] == ac.COMMON and (row["best_anyone"] is None or v > _value(row["best_anyone"])):
            row["best_anyone"] = e
        if row["best_any"] is None or v > _value(row["best_any"]):
            row["best_any"] = e
    out = []
    for row in by_key.values():
        cells = {k: [item_text(e, k) for e in sorted(items, key=lambda x: -(_value(x) or 0))]
                 for k, items in row["cells"].items()}
        best = row["best_anyone"]
        best_card = row["best_any"] if row["best_any"] is not best else None
        if any(cells.values()):                       # e.g. no automatic rate but a telco code
            row["notes"] = [n for n in row["notes"] if n != "No discount offered"]
        out.append({
            "market": row["market"], "route": row["route"], "ota": row["ota"], "airline": row["airline"],
            "status": row["status"], "travel_dates": ac._fmt_dates(row["dates"]), "seen": row["seen"],
            "basis": row["basis"], "cells": cells,
            "best_anyone": item_text(best, AUTO) if best else "",
            "best_with_card": item_text(best_card, BANK) if best_card else "",
            "notes": "; ".join(dict.fromkeys(row["notes"]))})
    ota_order = {lab: i for i, (lab, _k) in enumerate(ac.g.ROW_ORDER)}
    out.sort(key=lambda r: (r["status"] != "", r["market"] != "DOM", r["route"] in ("", ac.ALL_ROUTES),
                            r["route"], ota_order.get(r["ota"], 99), r["airline"]))
    return out


# --- Excel ------------------------------------------------------------------------------

_HEAD = [("market", "Market", 12), ("route", "Route", 10), ("travel_dates", "Travel date(s)", 13),
         ("ota", "OTA", 14), ("airline", "Airline", 8)] + \
        [(k, k, 24 if k != BANK else 34) for k in KINDS] + \
        [("best_anyone", "Best for anyone", 20), ("best_with_card", "Best with card / telco", 30),
         ("basis", "% of", 12), ("status", "Coverage", 13), ("notes", "Notes", 30)]
_NOTE = ("One row per route, OTA and airline; each column is a TYPE of discount, so the "
         "OTAs line up. A cell lists that type's promo codes, best first, as the % they were "
         "worth at the fares seen (caps applied; a range when fares differed). Automatic = no "
         "code needed. Telco on ShareTrip needs no verification (open to anyone); on FirstTrip "
         "the number is verified by OTP. % of: ShareTrip/FirstTrip base fare, GoZayaan booking "
         "total. Coverage 'Not searched' = rates from FirstTrip's coupon table for a route this "
         "run didn't search. Caps, fees and offer text: the (codes detail) sheet.")


def write_sheet(ws, report: dict[str, Any]) -> None:
    """The (all codes) sheet: one row per route x OTA x airline, discount types as columns."""
    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    st = ac.g._detail_styles()
    wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
    grey = PatternFill("solid", fgColor="EDEDED")
    t = ws.cell(1, 1, f"{report['report_date']} / {report['report_time']}hrs — all discounts by type")
    t.font = st["title"]
    ws.cell(2, 1, _NOTE).font = st["note"]
    for ci, (_key, head, width) in enumerate(_HEAD, start=1):
        c = ws.cell(3, ci, head)
        c.font, c.fill, c.alignment, c.border = st["head"], st["hdr_fill"], st["center"], st["border"]
        ws.column_dimensions[get_column_letter(ci)].width = width
    r = 4
    for row in rows(report.get("all_codes") or []):
        vals = {**row, **{k: "\n".join(v) for k, v in row["cells"].items()},
                "market": "Domestic" if row["market"] == "DOM" else "International"}
        for ci, (key, _h, _w) in enumerate(_HEAD, start=1):
            v = vals.get(key)
            cell = ws.cell(r, ci, v if v not in ("", None) else None)
            cell.font = st["label"] if key in ("best_anyone",) else st["data"]
            cell.alignment = wrap
            cell.border = st["border"]
            if row["status"]:
                cell.fill = grey
        r += 1
    ws.auto_filter.ref = f"A3:{get_column_letter(len(_HEAD))}{max(r - 1, 3)}"
    ws.freeze_panes = "F4"            # route/OTA/airline stay in view while scrolling types
    ws.print_title_rows = "3:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
