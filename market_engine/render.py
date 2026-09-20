"""Excel output for the schedule and fare views.

openpyxl only (already bundled for the discount grid); pandas/xlsxwriter are
excluded from the desktop build.

The sheets carry their own caveats — refused sources, held-back itineraries,
derived bases, data age — because a workbook gets mailed around long after the
screen that explained it is gone.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
NOTE_FONT = Font(color="9C6500", size=9, italic=True)
BEST_FILL = PatternFill("solid", fgColor="C6EFCE")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def _header(ws, row: int, labels: Iterable[str]) -> int:
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row, i, text)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center")
    return row + 1


def _widths(ws, widths: Iterable[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_schedule(ws, patterns, sched, title: str) -> None:
    ws.cell(1, 1, title).font = Font(bold=True, size=12)
    r = 2
    notes = []
    if sched.itineraries:
        notes.append(f"{sched.itineraries} connecting itineraries held back "
                     "(a marketed origin-destination is not an operated leg)")
    if sched.dropped_no_clock:
        notes.append(f"{sched.dropped_no_clock} rows excluded — no real departure time")
    if sched.dropped_no_flight:
        notes.append(f"{sched.dropped_no_flight} rows excluded — no readable flight number")
    for src, why in (sched.sources_refused or {}).items():
        notes.append(f"source refused: {src} — {why}")
    for n in notes:
        ws.cell(r, 1, n).font = NOTE_FONT
        r += 1
    r += 1

    r = _header(ws, r, ["Route", "Airline", "Flight", "Operates", "Dep", "Arr",
                        "Dates", "Sources", "Flags"])
    for p in patterns:
        flags = []
        if p.varies:
            flags.append("weekday pattern changes in range")
        if p.time_varies:
            flags.append("retimes by day")
        if p.disagreement:
            flags.append("sources disagree on time")
        ws.cell(r, 1, p.route)
        ws.cell(r, 2, p.airline)
        ws.cell(r, 3, p.flight)
        ws.cell(r, 4, p.weekday_label)
        ws.cell(r, 5, p.departure)
        ws.cell(r, 6, p.arrival)
        ws.cell(r, 7, p.dates)
        ws.cell(r, 8, ", ".join(p.sources))
        c = ws.cell(r, 9, "; ".join(flags))
        if flags:
            c.fill = WARN_FILL
        r += 1
    _widths(ws, [10, 8, 8, 24, 7, 7, 7, 18, 34])
    ws.freeze_panes = ws.cell(r - len(patterns), 1)


def write_fares(ws, table, best_by_route: dict, title: str) -> None:
    ws.cell(1, 1, title).font = Font(bold=True, size=12)
    r = 2
    if table.gross_only:
        ws.cell(r, 1, "gross only (base not trustworthy from these sources): "
                      + ", ".join(table.gross_only)).font = NOTE_FONT
        r += 1
    for src, why in (table.refused or {}).items():
        ws.cell(r, 1, f"source refused: {src} — {why}").font = NOTE_FONT
        r += 1
    r += 1

    r = _header(ws, r, ["Route", "Airline", "Flights", "Nonstop", "Dates",
                        "Gross low", "Gross high", "Gross avg",
                        "Base low", "Base high", "Base avg", "Data age (h)", "Flags"])
    for c_ in table.cells:
        flags = []
        if not c_.base_available:
            flags.append("base unavailable")
        elif c_.base_derived:
            flags.append("base partly derived")
        if c_.age_hours > 24:
            flags.append(f"{c_.age_hours/24:.0f}d old")
        ws.cell(r, 1, c_.route)
        ws.cell(r, 2, c_.airline)
        ws.cell(r, 3, c_.flights)
        ws.cell(r, 4, c_.nonstop)
        ws.cell(r, 5, c_.dates)
        for i, v in enumerate((c_.gross_low, c_.gross_high, c_.gross_avg), start=6):
            ws.cell(r, i, round(v)).number_format = "#,##0"
        for i, v in enumerate((c_.base_low, c_.base_high, c_.base_avg), start=9):
            cell = ws.cell(r, i, round(v) if c_.base_available else "—")
            if c_.base_available:
                cell.number_format = "#,##0"
        ws.cell(r, 12, round(c_.age_hours, 1))
        fc = ws.cell(r, 13, "; ".join(flags))
        if flags:
            fc.fill = WARN_FILL
        if best_by_route.get(c_.route) == c_.airline:
            for i in range(1, 14):
                if not ws.cell(r, i).fill.fgColor.rgb or ws.cell(r, i).fill.patternType is None:
                    ws.cell(r, i).fill = BEST_FILL
        r += 1
    _widths(ws, [10, 8, 8, 9, 7, 11, 11, 11, 10, 10, 10, 13, 30])


def write_timetable(ws, rows, title: str) -> None:
    """The operator's schedule sheet: sector-grouped, day flags, HHMM times.

    Sector is merged down its block because it labels a market, not a row, and
    both directions of that market sit inside it. Unknown seat capacity is left
    EMPTY rather than filled with a guess -- every load factor built on this
    column would inherit the error.
    """
    from market_engine.timetable import DAY_COLUMNS

    ws.cell(1, 1, title).font = Font(bold=True, size=12)
    header = (["Sector"] + list(DAY_COLUMNS)
              + ["Freq.", "Airline", "Aircraft", "Flight no.", "Org", "1Stop",
                 "Dest", "Dep", "Arr", "Seat Capacity", "Operated by"])
    r = _header(ws, 3, header)

    first_row = r
    for row in rows:
        ws.cell(r, 1, row.sector)
        for i, on in enumerate(row.days):
            c = ws.cell(r, 2 + i, 1 if on else None)
            c.alignment = Alignment(horizontal="center")
        ws.cell(r, 9, row.freq).alignment = Alignment(horizontal="center")
        ws.cell(r, 10, row.airline)
        ws.cell(r, 11, row.aircraft)
        ws.cell(r, 12, row.flight_no)
        ws.cell(r, 13, row.org).alignment = Alignment(horizontal="center")
        ws.cell(r, 14, row.stop or None).alignment = Alignment(horizontal="center")
        ws.cell(r, 15, row.dest).alignment = Alignment(horizontal="center")
        # Times are labels, not numbers: 0745 must keep its leading zero.
        for col, value in ((16, row.dep), (17, row.arr)):
            c = ws.cell(r, col, value or None)
            c.alignment = Alignment(horizontal="center")
            c.number_format = "@"
        cap = ws.cell(r, 18, row.seats if row.seats else None)
        cap.alignment = Alignment(horizontal="center")
        if not row.seats:
            cap.fill = WARN_FILL          # visibly unknown, not silently zero
        # Only a codeshare fills this: naming BG as BG's operator is noise.
        if row.operated_by:
            ws.cell(r, 19, "{} ({})".format(row.operated_by_name or row.operated_by,
                                            row.operated_by))
        r += 1

    # Merge each sector block so the sheet reads as markets, like the original.
    start = first_row
    for i in range(first_row, r + 1):
        at_end = i == r
        changed = (not at_end
                   and ws.cell(i, 1).value != ws.cell(start, 1).value)
        if at_end or changed:
            if i - start > 1:
                ws.merge_cells(start_row=start, start_column=1,
                               end_row=i - 1, end_column=1)
                ws.cell(start, 1).alignment = Alignment(horizontal="center",
                                                        vertical="center")
            start = i
    _widths(ws, [10, 4, 4, 4, 4, 4, 4, 4, 7, 22, 18, 11, 6, 7, 6, 7, 7, 14, 22])
    ws.freeze_panes = ws.cell(first_row, 1)


def write_workbook(path: Path, *, patterns=None, sched=None, table=None,
                   best_by_route: Optional[dict] = None, timetable=None,
                   title: str = "") -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    stamp = title or datetime.now().strftime("%d %b %Y %H:%M")
    if sched is not None:
        write_schedule(wb.create_sheet("Schedule"), patterns or [], sched,
                       f"Airline schedule — {stamp}")
    if timetable is not None:
        write_timetable(wb.create_sheet("Timetable"), timetable,
                        f"Schedule timetable — {stamp}")
    if table is not None:
        write_fares(wb.create_sheet("Fares"), table, best_by_route or {},
                    f"Fare comparison — {stamp}")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path
