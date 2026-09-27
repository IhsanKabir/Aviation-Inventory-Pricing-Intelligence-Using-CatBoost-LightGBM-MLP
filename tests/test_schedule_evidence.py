"""Schedule / timetable honesty: an OTA drops a SOLD-OUT flight, so a day it was
not seen is not proof it doesn't operate (field case 2026-09-27: VQ921 lost its
Friday), and one flight must not be split into rows by how sources spell its
aircraft."""
from __future__ import annotations

import json
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from market_engine.cache import CacheStore  # noqa: E402
from market_engine.collect import CollectResult  # noqa: E402
from market_engine.rows import clean_aircraft, from_connector  # noqa: E402
from market_engine.schedule import (build_schedule, searched_dates,  # noqa: E402
                                    weekday_states, weekly_pattern)
from market_engine.timetable import build_timetable  # noqa: E402

D0 = date(2026, 9, 28)                      # a Monday


def _offer(day: date, number: str, dep: str = "07:15", arr: str = "08:20",
           airline: str = "VQ", aircraft: str = "ATR725", stops: int = 0, via: str = ""):
    return from_connector({
        "airline": airline, "flight_number": number, "origin": "DAC", "destination": "CXB",
        "departure": f"{day.isoformat()}T{dep}:00", "arrival": f"{day.isoformat()}T{arr}:00",
        "stops": stops, "via_airports": via, "price_total_bdt": 4349, "fare_amount": 3224,
        "aircraft": aircraft}, source="FirstTrip", cabin="Economy")


def _days(n: int):
    return [D0 + timedelta(days=i) for i in range(n)]


def _timetable(rows, d1, d2):
    sched = build_schedule(CollectResult(rows=rows), date_from=d1, date_to=d2,
                           include_itineraries=True)
    return build_timetable(sched, searched=searched_dates(rows))


def test_the_rule_seen_unsure_not_operating_not_searched():
    fri = lambda n: {D0 + timedelta(days=4 + 7 * i) for i in range(n)}     # Fridays
    assert weekday_states({0, 1}, fri(1))[4] == "?"      # one Friday searched, not seen
    assert weekday_states({0, 1}, fri(2))[4] == ""       # two Fridays, never seen: a no
    assert weekday_states({0, 1}, set())[4] == "-"       # no Friday searched at all
    assert weekday_states({4}, fri(1))[4] == "1"
    assert weekday_states({0}, None)[4] == ""            # no search record: old behaviour


def test_sold_out_on_the_only_friday_is_unknown_not_a_missing_day():
    days = _days(10)                                   # 28 Sep..7 Oct: one Friday (2 Oct)
    rows = [_offer(d, "VQ 921") for d in days if d.weekday() != 4]
    rows += [_offer(d, "VQ 923", "10:30", "11:35") for d in days]  # the route WAS searched Fri
    vq921 = next(t for t in _timetable(rows, days[0], days[-1]) if t.flight_no == "VQ921")
    assert vq921.day_state == ("1", "1", "1", "1", "?", "1", "1")
    assert vq921.freq == 6                              # confirmed days only


def test_missing_on_every_one_of_several_fridays_is_a_real_no():
    days = _days(21)                                   # three Fridays
    rows = [_offer(d, "VQ 921") for d in days if d.weekday() != 4]
    rows += [_offer(d, "VQ 923", "10:30", "11:35") for d in days]
    vq921 = next(t for t in _timetable(rows, days[0], days[-1]) if t.flight_no == "VQ921")
    assert vq921.day_state[4] == ""


def test_weekly_label_names_the_unconfirmed_day():
    days = _days(10)
    rows = [_offer(d, "VQ 921") for d in days if d.weekday() != 4]
    rows += [_offer(d, "VQ 923", "10:30", "11:35") for d in days]
    sched = build_schedule(CollectResult(rows=rows), date_from=days[0], date_to=days[-1])
    pats = {p.flight: p for p in weekly_pattern(sched, date_from=days[0], date_to=days[-1],
                                                searched=searched_dates(rows))}
    assert pats["921"].weekday_label == "Mon/Tue/Wed/Thu/Sat/Sun (Fri?)"
    assert pats["923"].weekday_label == "Daily"


def test_one_aircraft_spelled_four_ways_is_one_row():
    days = _days(7)
    spellings = ["ATR725", "ATR 72-500", "", "ATR TURBOPROP", "ATR725", "ATR725", ""]
    rows = [_offer(d, "VQ 921", aircraft=a) for d, a in zip(days, spellings)]
    tt = [t for t in _timetable(rows, days[0], days[-1]) if t.flight_no == "VQ921"]
    assert len(tt) == 1 and tt[0].freq == 7 and tt[0].aircraft == "ATR725"
    assert tt[0].seats == 72


def test_dash8_names_from_different_sources_merge():
    days = _days(4)
    rows = [_offer(d, "BG 433", "10:15", "11:30", "BG", a)
            for d, a in zip(days, ["DH8", "DEHAVILLAND DASH 8", "DH8", "Dash 8-Q400"])]
    tt = [t for t in _timetable(rows, days[0], days[-1]) if t.flight_no == "BG433"]
    assert len(tt) == 1 and tt[0].aircraft == "DH8"


def test_genuinely_different_aircraft_stay_separate():
    days = _days(4)
    rows = [_offer(d, "BG 202", "10:05", "11:00", "BG", a)
            for d, a in zip(days, ["Boeing 787-8", "Boeing 787-9", "Boeing 787-8", "Boeing 787-9"])]
    tt = [t for t in _timetable(rows, days[0], days[-1]) if t.flight_no == "BG202"]
    assert sorted(t.aircraft for t in tt) == ["Boeing 787-8", "Boeing 787-9"]


def test_a_counted_but_unnamed_stop_is_not_shown_as_nonstop():
    days = _days(2)
    rows = [_offer(days[0], "BS 322", "02:00", "08:50", "BS"),
            _offer(days[1], "BS 322", "02:00", "10:50", "BS", stops=1)]
    stops = {t.arr: t.stop for t in _timetable(rows, days[0], days[-1])}
    assert stops == {"0850": "", "1050": "1 stop"}


def test_sharetrip_aircraft_record_becomes_its_model():
    assert clean_aircraft({"code": "725", "model": "ATR 72-500 "}) == "ATR 72-500"
    assert clean_aircraft("{'code': '725', 'model': 'ATR 72-500 '}") == "ATR 72-500"  # old cache text
    assert clean_aircraft("Boeing 737-800") == "Boeing 737-800"
    assert clean_aircraft(None) == ""


def test_old_cache_rows_are_repaired_on_read():
    store = CacheStore(Path(tempfile.mkdtemp()) / "m.sqlite")
    store.put("ShareTrip", "DAC", "CXB", "2026-10-02", "Economy", [_offer(date(2026, 10, 2), "VQ 921")])
    import sqlite3
    con = sqlite3.connect(store.path if hasattr(store, "path") else store._path)
    k, blob = con.execute("select k, payload from q").fetchone()
    rows = json.loads(blob)
    rows[0]["aircraft"] = "{'code': '725', 'model': 'ATR 72-500 '}"     # as v0.2.x stored it
    del rows[0]["operated_by"]                                           # pre-codeshare rows
    con.execute("update q set payload=? where k=?", (json.dumps(rows), k))
    con.commit()
    got = store.get("ShareTrip", "DAC", "CXB", "2026-10-02", "Economy", max_age=timedelta(days=9999))
    assert got[0].aircraft == "ATR 72-500" and got[0].operated_by == ""


def test_merged_sources_keep_the_most_specific_aircraft():
    """FirstTrip and ShareTrip list the same flight on the same date; the leg must
    keep the model number, whichever source was read first."""
    day = date(2026, 10, 1)
    sharetrip = _offer(day, "BS 141", "07:15", "08:20", "BS", "ATR TURBOPROP")
    firsttrip = _offer(day, "BS 141", "07:15", "08:20", "BS", "ATR 72 - 600")
    for order in ([sharetrip, firsttrip], [firsttrip, sharetrip]):
        tt = [t for t in _timetable(order, day, day) if t.flight_no == "BS141"]
        assert len(tt) == 1 and tt[0].aircraft == "ATR 72 - 600" and tt[0].seats == 70


def test_dash8_under_any_name_gets_its_seat_count():
    from market_engine.timetable import seat_capacity
    table = {"BG": {"DH8": 74}}
    assert seat_capacity("BG", "DEHAVILLAND DASH 8", table) == 74
    assert seat_capacity("BG", "Dash 8-Q400", table) == 74
    assert seat_capacity("BG", "Boeing 787-9", table) is None


# --- settling "?" from the same weekday in nearby weeks --------------------------

def test_iata_season_boundary():
    from market_engine.confirm import season
    assert season(date(2026, 10, 24)) == ("S", 2026)      # last Sunday of Oct 2026 = 25th
    assert season(date(2026, 10, 25)) == ("W", 2026)
    assert season(date(2027, 3, 27)) == ("W", 2026)
    assert season(date(2027, 3, 28)) == ("S", 2027)


def test_unsure_targets_look_at_the_same_weekday_in_the_same_season():
    from market_engine.confirm import unsure_targets
    days = _days(10)
    rows = [_offer(d, "VQ 921") for d in days if d.weekday() != 4]
    rows += [_offer(d, "VQ 923", "10:30", "11:35") for d in days]
    tt = _timetable(rows, days[0], days[-1])
    targets = unsure_targets(tt, days[0], days[-1], today=date(2026, 9, 27))
    assert targets == {("DAC", "CXB"): [date(2026, 10, 9), date(2026, 10, 16)]}
    # Past dates are never targets, and the winter season (from 25 Oct) is not mixed in.
    late = unsure_targets(tt, days[0], days[-1], today=date(2026, 10, 10))
    assert late == {("DAC", "CXB"): [date(2026, 10, 16)]}


def _admin_api(monkeypatch, cached: dict):
    """DesktopApi for the admin whose collect() answers only from `cached`
    ({date: rows}) - no network - and records every date it was asked for."""
    import market_engine.collect as mc
    from desktop import backend as bm
    asked = []

    def fake_collect(plan, cache=None, progress=None, should_cancel=None):
        rows = []
        for day in plan.dates if plan.source_keys else []:   # like collect(): no source, no query
            asked.append(day)
            rows += cached.get(day, [])
        return mc.CollectResult(rows=rows)
    monkeypatch.setattr(mc, "collect", fake_collect)
    cfg = Path(tempfile.mkdtemp())
    monkeypatch.setattr(bm, "config_dir", lambda: cfg)
    monkeypatch.setattr(bm, "_HAS_KEYRING", False)
    api = bm.DesktopApi()
    api._store_token("tok")
    api._config.update(is_admin=True, har_dir=str(cfg))
    monkeypatch.setattr(api, "check_access",
                        lambda: {"status": "approved", "allowed": True, "is_admin": True})
    monkeypatch.setattr(api, "_log_usage", lambda *a, **k: None)
    return api, asked


def _week_rows(days, friday_921: bool):
    rows = {}
    for d in days:
        day_rows = [_offer(d, "VQ 923", "10:30", "11:35")]
        if d.weekday() != 4 or friday_921:
            day_rows.append(_offer(d, "VQ 921"))
        rows[d] = day_rows
    return rows


def test_sold_out_friday_is_settled_as_flying(monkeypatch):
    rng = _days(10)                                      # 28 Sep..7 Oct, VQ921 sold out 2 Oct
    cached = _week_rows(rng, friday_921=False)
    cached.update(_week_rows([date(2026, 10, 9), date(2026, 10, 16)], friday_921=True))
    api, asked = _admin_api(monkeypatch, cached)
    r = api.run_market("schedule", "DAC-CXB", "2026-09-28", "2026-10-07", ["firsttrip"])
    vq = next(t for t in api._market_last["timetable"] if t.flight_no == "VQ921")
    assert vq.day_state == ("1",) * 7 and vq.freq == 7
    assert "VQ921 DAC-CXB 0715: flies Fri" in r["confirmed"]
    assert {d for d in asked if d > date(2026, 10, 7)} == {date(2026, 10, 9), date(2026, 10, 16)}


def test_a_friday_missing_every_week_is_settled_as_no_flight(monkeypatch):
    rng = _days(10)
    cached = _week_rows(rng, friday_921=False)
    cached.update(_week_rows([date(2026, 10, 9), date(2026, 10, 16)], friday_921=False))
    api, _ = _admin_api(monkeypatch, cached)
    r = api.run_market("schedule", "DAC-CXB", "2026-09-28", "2026-10-07", ["firsttrip"])
    vq = next(t for t in api._market_last["timetable"] if t.flight_no == "VQ921")
    assert vq.day_state[4] == "" and vq.freq == 6
    assert "VQ921 DAC-CXB 0715: no flight Fri" in r["confirmed"]


def test_extra_live_lookups_are_capped(monkeypatch):
    rng = _days(10)
    api, asked = _admin_api(monkeypatch, _week_rows(rng, friday_921=False))
    monkeypatch.setattr(type(api), "CONFIRM_MAX_QUERIES", 1)
    api.run_market("schedule", "DAC-CXB", "2026-09-28", "2026-10-07", ["firsttrip"])
    assert len([d for d in asked if not rng[0] <= d <= rng[-1]]) == 1   # nearest week only


def test_har_only_users_settle_from_their_other_capture_dates(monkeypatch):
    from market_engine import har as har_mod
    rng = _days(10)
    rows = [r for rs in _week_rows(rng, friday_921=False).values() for r in rs]
    rows += _week_rows([date(2026, 10, 9)], friday_921=True)[date(2026, 10, 9)]
    api, asked = _admin_api(monkeypatch, {})
    api._config["is_admin"] = False
    monkeypatch.setattr(api, "check_access",
                        lambda: {"status": "approved", "allowed": True, "is_admin": False})
    monkeypatch.setattr(har_mod, "collect_har_rows",
                        lambda d, purpose="fare", progress=None, channels=None: (list(rows), []))
    r = api.run_market("schedule", "DAC-CXB", "2026-09-28", "2026-10-07", ["har"])
    vq = next(t for t in api._market_last["timetable"] if t.flight_no == "VQ921")
    assert vq.day_state[4] == "1" and asked == []                     # no live lookups
    assert r["rows"] == len([x for x in rows if x.departure_date <= date(2026, 10, 7)])


def test_timetable_sheet_explains_question_marks():
    from openpyxl import Workbook
    from market_engine.render import write_timetable
    days = _days(10)
    rows = [_offer(d, "VQ 921") for d in days if d.weekday() != 4]
    rows += [_offer(d, "VQ 923", "10:30", "11:35") for d in days]
    ws = Workbook().active
    write_timetable(ws, _timetable(rows, days[0], days[-1]), "Schedule timetable")
    assert "? = not on sale on the ONLY such day searched" in ws.cell(2, 1).value
    values = [[c.value for c in r] for r in ws.iter_rows(min_row=4)]
    vq921 = next(v for v in values if v[11] == "VQ921")
    assert vq921[1:8] == [1, 1, 1, 1, "?", 1, 1] and vq921[8] == 6
