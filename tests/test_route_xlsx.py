"""Route Excel sheets + blank-row removal across the workbook."""
from __future__ import annotations

import tempfile
from pathlib import Path

from openpyxl import load_workbook

from discount_engine import by_route, grid, route_xlsx


def _block(route: str, cells_by_label: dict) -> dict:
    o, d = route.split("-")
    return by_route.highlighted(by_route.route_blocks({(o, d): cells_by_label}))[0]


def test_blank_rows_and_a_dangling_separator_are_dropped():
    rows = [{"label": "USBA OTA B2B", "kind": "b2b", "cells": {"BS": ""}},
            {"label": "BDFare", "kind": "b2b", "cells": {"BS": "7"}},
            {"label": "__sep__", "kind": "sep", "cells": {}},
            {"label": "Amy", "kind": "b2c", "cells": {"BS": " "}}]
    assert [r["label"] for r in route_xlsx.visible_rows(rows)] == ["BDFare"]


def test_coupon_tiers_go_on_their_own_lines():
    assert route_xlsx.cell_lines("9.5(Bkash, 2% fee), 18 (Stellar Signature, 2% fee)") == [
        "9.5% (Bkash, 2% fee)", "18% (Stellar Signature, 2% fee)"]
    assert route_xlsx.cell_lines("~7.47") == ["~7.47%"]
    assert route_xlsx.cell_lines("") == []


def test_overview_keeps_b2b_and_b2c_apart_and_never_crowns_zero():
    block = _block("DAC-DXB", {"BDFare": {"BS": "~7.5", "MH": "0"},
                               "ShareTrip-B2C": {"BS": "9(Bkash, 2% fee)"}})
    lines, top = route_xlsx._best_lines(block, "BS")
    assert lines == ["B2C 9% · ST-B2C", "B2B ~7.5% · BDFare"]      # both, labelled
    assert top == 7.5                                               # 9 - 2% fee = 7 net
    lines, top = route_xlsx._best_lines(block, "MH")
    assert lines == ["B2B no discount · BDFare"] and top == 0


def _report(**extra) -> dict:
    blocks = [_block("DAC-CGP", {"BDFare": {"BS": "7", "VQ": "7"}, "Amy": {"BS": "4"}}),
              _block("DAC-DXB", {"AKIJ AIR-B2B": {"EK": "6"}})]
    grids = {"DOM": {"columns": ["BS"], "rows": [
        {"label": lab, "kind": kind, "cells": {"BS": "7" if lab == "BDFare" else ""}}
        for lab, kind in grid.ROW_ORDER]}}
    return {"report_date": "20/09/2026", "report_time": "1800", "grids": grids,
            "by_route": blocks, **extra}


def _write(report: dict):
    return load_workbook(grid.write_single_sheet_xlsx(report, None,
                                                      Path(tempfile.mkdtemp()) / "r.xlsx"))


def test_daily_sheet_drops_blank_rows_but_names_failed_captures():
    wb = _write(_report(channel_status={"ShareTrip-B2C": "captured_but_empty", "BDFare": "ok"}))
    col_a = [c.value for c in wb["20 September"]["A"] if c.value]
    assert "BDFare" in col_a and "TLN" not in col_a and "SHARETRIP-B2B" not in col_a
    assert "Captured but no data (check these captures): ShareTrip-B2C" in col_a[0]
    by_airline = [c.value for c in wb["20 September (by airline)"][2] if c.value]
    assert by_airline == ["Airline", "BDFare", "Best OTA"]         # blank OTA columns gone


def test_coverage_counts_airlines_not_just_presence():
    ws = _write(_report())["20 September routes"]
    rows = {r[0]: r[1:] for r in ws.iter_rows(values_only=True) if r[0] in ("DAC-CGP", "DAC-DXB")}
    heads = next(r for r in ws.iter_rows(values_only=True) if r[0] == "Route" and "Amy" in r)
    cov = dict(zip(heads[1:], rows["DAC-CGP"]))
    assert cov["BDFare"] == "2/2" and cov["Amy"] == "1/2" and cov["AKIJ AIR-B2B"] == "–"


def test_route_links_point_at_their_detail_block():
    wb = _write(_report())
    link = next(c for c in wb["20 September routes"]["A"] if c.value == "DAC-DXB").hyperlink
    target_row = int(link.location.split("!A")[1]) if link.location else int(link.target.split("!A")[1])
    assert wb["20 September route grids"].cell(target_row, 1).value.startswith("DAC-DXB")


def test_detail_by_route_groups_airlines_under_each_route():
    report = by_route.with_preferred(_report(), ["DAC-DXB", "DAC-JED"])
    ws = _write(report)["20 September (detail by route)"]
    col_a = [c.value for c in ws["A"] if c.value]
    assert col_a[:2] == ["20/09/2026 / 1800hrs — detailed by route", "Airline"]
    order = ["★ Your preferred routes", "Not captured yet: DAC-JED",
             "DAC-DXB · International  ★", "EK",
             "Other routes found in this run", "DAC-CGP · Domestic", "BS", "VQ"]
    assert [v for v in col_a if v in order] == order          # route band, then its airlines
    bs_otas = [r[1] for r in ws.iter_rows(values_only=True) if r[1] in ("BDFare", "Amy")]
    assert bs_otas[:2] == ["BDFare", "Amy"]                   # BS on CGP: both OTAs, one row each


def test_every_sheet_name_fits_excels_limit_even_in_september():
    wb = _write(by_route.with_preferred(_report(), ["DAC-CGP"]))
    assert all(len(n) <= 31 for n in wb.sheetnames)


def test_data_sheet_has_one_filterable_row_per_rate():
    ws = _write(_report())["20 September route data"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 4                          # CGP: BDFare BS+VQ, Amy BS; DXB: AKIJ EK
    assert ws.auto_filter.ref.startswith("A1:")
    first = dict(zip(route_xlsx.DATA_HEADERS, rows[0]))
    assert first["Route"] == "DAC-CGP" and first["Type"] == "B2B" and first["Rate %"] == 0.07
