"""Per-route discount view: same cell rules per route; the summary never changes."""
from __future__ import annotations

import json
import tempfile
from pathlib import Path

import pytest

from discount_engine import by_route, grid
from discount_engine.sanitize import sanitize_report_for_sync
from modules import gozayaan_har, sharetrip_har


def _har(entries: list) -> str:
    p = Path(tempfile.mkdtemp()) / "cap.har"
    p.write_text(json.dumps({"log": {"entries": entries}}), encoding="utf-8")
    return str(p)


def _entry(url: str, resp: dict, body: dict | None = None) -> dict:
    req = {"method": "POST" if body is not None else "GET", "url": url}
    if body is not None:
        req["postData"] = {"text": json.dumps(body)}
    return {"request": req, "response": {"content": {"text": json.dumps(resp)}}}


# --- ShareTrip: route from the itinerary legs, de-duplicated per route ----------------

def _st_flight(airline: str, origin: str, dest: str, pct: float, base: int) -> dict:
    return {"legs": [{"marketingAirline": airline, "origin": {"code": origin},
                      "destination": {"code": dest}}],
            "domestic": False, "displayPrice": {"discount": pct, "totalFare": {"base": base}}}


def test_sharetrip_search_rows_keep_their_route_and_both_routes_survive():
    url = "https://api.sharetrip.net/api/v2/flight/search/available-flights?searchId=1"
    har = _har([_entry(url, {"response": {"matchedFlights": [
        _st_flight("BS", "DAC", "DXB", 7.1, 40000),
        _st_flight("BS", "DAC", "KUL", 7.1, 30000)]}})])
    routed = sharetrip_har.parse_discounts_routed(har)
    assert {(r["origin"], r["destination"]) for r in routed} == {("DAC", "DXB"), ("DAC", "KUL")}
    # the plain parser (summary input) still de-duplicates across the file, unchanged
    assert len(sharetrip_har.parse_discounts(har)) == 1


# --- GoZayaan: route joined from the search that produced the coupon request ----------

def _gz_search(sid: str, origin: str, dest: str) -> dict:
    return _entry("https://production.gozayaan.com/api/flight/v4.0/search/",
                  {"result": {"search_id": sid}},
                  {"trips": [{"origin": origin, "destination": dest}]})


def _gz_coupons(sid: str, carrier: str, price: int) -> dict:
    coupon = {"discount_promo_code": "X7", "discount_markup": {
        "markup_type": "PERCENTAGE", "markup_amount": 7, "markup_max_amount": 0}}
    return _entry("https://production.gozayaan.com/api/business_rules/get_discount_list/",
                  {"result": [coupon]},
                  {"search_id": sid, "plating_carrier": carrier, "flight_type": "DOM",
                   "product_price": price})


def test_gozayaan_coupon_gets_its_searchs_route_on_every_route():
    har = _har([_gz_search("a", "DAC", "CXB"), _gz_coupons("a", "BS", 5000),
                _gz_search("b", "DAC", "CGP"), _gz_coupons("b", "BS", 4500),
                _gz_coupons("zz", "BG", 5000)])          # search not captured
    rows = gozayaan_har.parse_discounts_routed(har)
    assert sorted((r["origin"], r["destination"], r["airline"]) for r in rows) == [
        ("", "", "BG"), ("DAC", "CGP", "BS"), ("DAC", "CXB", "BS")]
    assert len(gozayaan_har.parse_discounts(har)) == 2   # summary input unchanged


# --- blocks: ordering, columns, coverage -----------------------------------------------

def test_blocks_put_domestic_first_and_name_missing_otas():
    table = {("DAC", "DXB"): {"BDFare": {"EK": "5", "BS": "7"}},
             ("DAC", "CXB"): {"Amy": {"VQ": "4", "BS": "6"}}}
    blocks = by_route.route_blocks(table)
    assert [b["route"] for b in blocks] == ["DAC-CXB", "DAC-DXB"]
    assert blocks[0]["market"] == "DOM" and blocks[0]["columns"] == ["BS", "VQ"]
    assert blocks[1]["columns"] == ["BS", "EK"]                    # report column order
    assert blocks[0]["coverage"]["with_data"] == ["Amy"]
    assert "BDFare" in blocks[0]["coverage"]["without"]
    colored = by_route.highlighted(blocks)
    assert colored[0]["best"]["BS"]["channel"] == "Amy"


# --- end to end through build_report ---------------------------------------------------

def _bd_row(airline: str, dest: str, gross: int, agent: int, pct: float) -> dict:
    return {"channel": "bdfare", "persona": "B2B", "airline": airline,
            "origin": "DAC", "destination": dest, "domestic": False,
            "gross_bdt": gross, "agent_bdt": agent, "customer_net_bdt": None,
            "base_est_bdt": gross, "base_source": "exact",
            "commission_bdt": gross - agent, "commission_pct": pct}


def test_build_report_splits_by_route_without_touching_the_summary(monkeypatch):
    rows = [_bd_row("EK", "DXB", 60000, 57000, 5.0),      # cheapest EK overall -> summary
            _bd_row("EK", "JED", 80000, 74000, 7.5)]
    monkeypatch.setattr(grid.bdfare_har, "parse_commissions", lambda p, **kw: rows)
    har = _har([])
    report = grid.build_report(None, [], bdfare_hars=[har], use_true_base=False)

    intl = {r["label"]: r["cells"] for r in report["grids"]["INTL"]["rows"]}
    assert intl["BDFare"]["EK"] == "5"                     # summary: cheapest offer, as before
    per_route = {b["route"]: {r["label"]: r["cells"] for r in b["rows"]}
                 for b in report["by_route"]}
    assert per_route["DAC-DXB"]["BDFare"]["EK"] == "5"
    assert per_route["DAC-JED"]["BDFare"]["EK"] == "7.5"   # JED keeps its own rate
    assert grid._PARSE_MEMO is None                        # memo ends with the run


def test_each_har_is_parsed_once_per_run(monkeypatch):
    calls = {"n": 0}

    def parse(p, **kw):
        calls["n"] += 1
        return [_bd_row("EK", "DXB", 60000, 57000, 5.0)]
    monkeypatch.setattr(grid.bdfare_har, "parse_commissions", parse)
    grid.build_report(None, [], bdfare_hars=[_har([])], use_true_base=False)
    assert calls["n"] == 1          # summary + per-route share one parse


def test_memo_is_cleared_even_when_a_run_fails(monkeypatch):
    monkeypatch.setattr(grid, "_build_report", lambda *a, **k: 1 / 0)
    with pytest.raises(ZeroDivisionError):
        grid.build_report(None, [])
    assert grid._PARSE_MEMO is None


def test_excel_gets_a_route_sheet_only_when_routes_exist():
    from openpyxl import load_workbook
    blocks = by_route.route_blocks({("DAC", "CXB"): {"Amy": {"BS": "6"}},
                                    ("DAC", "DXB"): {"BDFare": {"EK": "5"}}})
    report = {"report_date": "20/09/2026", "report_time": "1800", "grids": {},
              "by_route": by_route.highlighted(blocks)}
    out = Path(tempfile.mkdtemp())
    wb = load_workbook(grid.write_single_sheet_xlsx(report, None, out / "a.xlsx"))
    assert wb.sheetnames[-4:] == ["20 September (detail by route)", "20 September routes",
                                  "20 September route grids", "20 September route data"]
    detail = [c.value for c in wb["20 September route grids"]["A"] if c.value]
    assert "DAC-CXB · Domestic" in detail
    assert "No data on this route: BDFare" in detail       # captured elsewhere, not here
    del report["by_route"]                                   # a stored/synced report
    wb = load_workbook(grid.write_single_sheet_xlsx(report, None, out / "b.xlsx"))
    assert not any("route" in n for n in wb.sheetnames)


def test_routes_never_reach_the_sync_payload():
    report = {"report_date": "20/09/2026", "grids": {}, "by_route": [
        {"route": "DAC-DXB", "market": "INTL", "columns": [], "rows": []}]}
    payload = sanitize_report_for_sync(report)
    assert "by_route" not in payload and "DAC-DXB" not in json.dumps(payload)
