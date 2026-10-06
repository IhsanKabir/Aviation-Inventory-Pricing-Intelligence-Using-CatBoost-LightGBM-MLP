"""All codes by type: one row per route x OTA x airline, discount types as columns."""
from __future__ import annotations

import tempfile
from pathlib import Path

from discount_engine import all_codes as ac, codes_by_type as cbt, grid


def _e(ota, airline, tier, code, who, eff, route="DAC-CXB", **kw):
    return ac._obs(market="DOM", route=route, ota=ota, airline=airline, tier=tier, code=code, who=who,
                   published=kw.pop("published", eff), effective=eff, basis=kw.pop("basis", "Base fare"),
                   date="2026-10-31", fare=(route, "2026-10-31", 1), **kw)


ENTRIES = ac.merge([
    _e("Firsttrip-B2C", "BS", "Common", "FTBSDOM", "Anyone (automatic)", 15.0),
    _e("Firsttrip-B2C", "BS", "Common", "FT-bKASH", "Pay with bKash", 12.0),
    _e("Firsttrip-B2C", "BS", "Common", "FT-Nagad", "Pay with Nagad", 12.0),
    _e("Firsttrip-B2C", "BS", "Special", "FTGPSTAR", "GP customers (number verified by OTP)", 14.0),
    _e("ShareTrip-B2C", "BS", "Common", "FLIGHTINT", "Anyone (automatic)", 8.5),
    _e("ShareTrip-B2C", "BS", "Common", "FLYGPSTAR", "Anyone entering the GPStar code (not verified)", 9.5),
    _e("ShareTrip-B2C", "BS", "Special", "STLRSIQ326", "Stellar Signature", 18.0),
    _e("ShareTrip-B2C", "BS", "No discount", "ZEROEMI", "ZEROEMI", None, published=0.0),
    _e("Go Zayaan", "BS", "Common", "RGOFLY", "Any card", 7.0, basis="Booking total"),
    _e("Go Zayaan", "BS", "Common", "RFLYDOM", "Pay with Nagad / Upay / Tap / Rocket", 7.0),
    _e("Go Zayaan", "BS", "Special", "RAMEX", "City Bank AMEX", 13.0),
    _e("Firsttrip-B2C", "BS", "Not searched", "FT-Nagad", "Pay with Nagad", None, route="DAC-ZYL",
       published=12.0),
])


def _row(rows, ota, route="DAC-CXB"):
    return next(r for r in rows if r["ota"] == ota and r["route"] == route)


def test_each_option_lands_in_its_type_column():
    rows = cbt.rows(ENTRIES)
    ft = _row(rows, "Firsttrip-B2C")
    assert ft["cells"]["Automatic"] == ["15% FTBSDOM"]
    assert ft["cells"]["bKash"] == ["12% FT-bKASH"]
    assert ft["cells"]["Nagad & other wallets"] == ["12% FT-Nagad"]
    assert ft["cells"]["Telco"] == ["14% FTGPSTAR (GP)"]
    st = _row(rows, "ShareTrip-B2C")
    assert st["cells"]["Telco"] == ["9.5% FLYGPSTAR (GPStar)"]           # open on ShareTrip
    assert st["cells"]["Bank & card offers"] == ["18% STLRSIQ326 (Stellar Signature)"]
    gz = _row(rows, "Go Zayaan")
    assert gz["cells"]["Any card / payment"] == ["7% RGOFLY"]
    assert gz["cells"]["Nagad & other wallets"] == ["7% RFLYDOM"]


def test_best_columns_and_zero_codes_left_out():
    rows = cbt.rows(ENTRIES)
    ft, st = _row(rows, "Firsttrip-B2C"), _row(rows, "ShareTrip-B2C")
    assert ft["best_anyone"] == "15% FTBSDOM" and ft["best_with_card"] == ""    # GP 14 < 15
    assert st["best_anyone"] == "9.5% FLYGPSTAR" and st["best_with_card"] == "18% STLRSIQ326 (Stellar Signature)"
    assert not any("ZEROEMI" in t for r in rows for items in r["cells"].values() for t in items)


def test_routes_not_searched_stay_their_own_rows():
    rows = cbt.rows(ENTRIES)
    zyl = _row(rows, "Firsttrip-B2C", route="DAC-ZYL")
    assert zyl["status"] == "Not searched" and zyl["cells"]["Nagad & other wallets"] == ["12% FT-Nagad"]
    assert rows[-1] is zyl                                     # listed after every searched row


def test_excel_all_codes_sheet_has_type_columns_and_a_detail_sheet():
    from openpyxl import load_workbook
    report = {"report_date": "04/10/2026", "report_time": "1200", "generated_at": "x",
              "grids": {"DOM": {"columns": [], "rows": []}, "INTL": {"columns": [], "rows": []}},
              "channel_status": {}, "sources": {}, "all_codes": ENTRIES}
    path = Path(tempfile.mkdtemp()) / "r.xlsx"
    grid.write_single_sheet_xlsx(report, None, path)
    wb = load_workbook(path)
    ws = wb["04 October (all codes)"]
    heads = [c.value for c in ws[3]]
    assert heads[5:11] == ["Automatic", "bKash", "Nagad & other wallets", "Telco", "Any card / payment",
                           "Bank & card offers"]
    first = {h: c.value for h, c in zip(heads, ws[4])}
    assert (first["Route"], first["OTA"], first["Automatic"], first["Best for anyone"]) == \
        ("DAC-CXB", "Firsttrip-B2C", "15% FTBSDOM", "15% FTBSDOM")
    assert "04 October (codes detail)" in wb.sheetnames


def test_empty_type_shows_zero_and_unknown_types_say_not_captured():
    rows = cbt.rows(ENTRIES)
    bg_like = _row(rows, "Go Zayaan")
    assert bg_like["empty"]["Automatic"] == "0" and bg_like["empty"]["Telco"] == "0"
    assert "Other" not in bg_like["empty"] or bg_like["empty"]["Other"] == ""
    # FirstTrip search only: its wallet/card/telco coupons weren't captured -> unknown
    rows = cbt.rows(ENTRIES, cbt.uncaptured_kinds({"ft_offer_list_missing": True}))
    ft = _row(rows, "Firsttrip-B2C")
    assert ft["empty"].get("Any card / payment") == "not captured"
    assert _row(rows, "Firsttrip-B2C", route="DAC-ZYL")["empty"]["Automatic"] == "not searched"


def test_airline_with_nothing_off_reads_zero_everywhere():
    vq = ac.merge([_e("Firsttrip-B2C", "VQ", "No discount", "(none)", "No discount on this fare", None,
                      published=0.0)])
    (row,) = cbt.rows(vq)
    assert row["best_anyone"] == "0" and row["empty"]["Automatic"] == "0" and row["empty"]["bKash"] == "0"
